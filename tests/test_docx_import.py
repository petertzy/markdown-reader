import os
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend.renderer import render_markdown
from backend.routers.files import _convert_docx_to_markdown


def _save(document) -> str:
    tmpdir = tempfile.mkdtemp()
    path = os.path.join(tmpdir, "sample.docx")
    document.save(path)
    return path


def _build_table(rows, cols, values):
    document = Document()
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    for row_index, row in enumerate(values):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    return document


class TestDocxTableImport(unittest.TestCase):
    def test_table_gets_a_gfm_delimiter_row(self):
        document = _build_table(2, 2, [["Name", "Age"], ["Bob", "30"]])

        markdown = _convert_docx_to_markdown(_save(document))

        # GFM needs `| --- |` under the header; without it the pipes are shown
        # as literal text instead of a table.
        self.assertIn("| Name | Age |", markdown)
        self.assertIn("| --- | --- |", markdown)
        self.assertIn("| Bob | 30 |", markdown)

    def test_imported_table_actually_renders_as_a_table(self):
        document = _build_table(2, 2, [["Name", "Age"], ["Bob", "30"]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertIn("<table>", render_markdown(markdown))

    def test_delimiter_row_matches_the_header_width(self):
        document = _build_table(2, 3, [["A", "B", "C"], ["1", "2", "3"]])

        markdown = _convert_docx_to_markdown(_save(document))

        lines = markdown.splitlines()
        self.assertEqual(
            len([cell for cell in lines[0].split("|") if cell.strip()]),
            len([cell for cell in lines[1].split("|") if cell.strip()]),
        )

    def test_horizontally_merged_cell_is_not_duplicated(self):
        document = _build_table(2, 2, [["Report", "Value"], ["A", "1"]])
        table = document.tables[0]
        # Merge the two header cells into one cell spanning both grid columns.
        kept, dropped = table.cell(0, 0)._tc, table.cell(0, 1)._tc
        dropped.getparent().remove(dropped)
        grid_span = OxmlElement("w:gridSpan")
        grid_span.set(qn("w:val"), "2")
        kept.get_or_add_tcPr().append(grid_span)

        markdown = _convert_docx_to_markdown(_save(document))

        # python-docx yields the merged cell once per spanned column, so the
        # header text used to appear twice.
        self.assertEqual(markdown.count("Report"), 1)
        self.assertNotIn("| Report | Report |", markdown)

    def test_merged_table_still_renders_as_a_table(self):
        document = _build_table(2, 2, [["Report", "Value"], ["A", "1"]])
        table = document.tables[0]
        kept, dropped = table.cell(0, 0)._tc, table.cell(0, 1)._tc
        dropped.getparent().remove(dropped)
        grid_span = OxmlElement("w:gridSpan")
        grid_span.set(qn("w:val"), "2")
        kept.get_or_add_tcPr().append(grid_span)

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertIn("<table>", render_markdown(markdown))

    def test_distinct_cells_with_identical_text_are_both_kept(self):
        # Identity dedupe must not degrade into a text dedupe.
        document = _build_table(1, 2, [["same", "same"]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertIn("| same | same |", markdown)

    def test_vertically_merged_cell_appears_once_per_spanned_row(self):
        # A vertically merged cell is yielded once per spanned row, carrying the
        # content of the cell it continues, so its text belongs on both rows.
        document = _build_table(3, 2, [["Group", "Total"], ["spans", "5"], ["", "7"]])
        table = document.tables[0]
        table.cell(1, 0).merge(table.cell(2, 0))

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertIn("| spans | 5 |", markdown)
        self.assertIn("| spans | 7 |", markdown)
        # Neither row duplicates the cell into two adjacent columns.
        self.assertNotIn("| spans spans |", markdown)

    def test_newlines_inside_a_cell_are_collapsed(self):
        document = _build_table(1, 1, [["line one\nline two"]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertIn("| line one line two |", markdown)

    def test_a_pipe_inside_a_cell_is_escaped(self):
        document = _build_table(2, 2, [["Name", "Note"], ["Bob", "3 | 4"]])

        markdown = _convert_docx_to_markdown(_save(document))

        # An unescaped pipe would split the cell into two extra columns.
        self.assertIn("| Bob | 3 \\| 4 |", markdown)
        self.assertNotIn("| Bob | 3 | 4 |", markdown)

    def test_escaped_pipe_cell_still_renders_as_one_table(self):
        document = _build_table(2, 2, [["Name", "Note"], ["Bob", "3 | 4"]])

        html = render_markdown(_convert_docx_to_markdown(_save(document)))

        self.assertIn("<td>3 | 4</td>", html)

    def test_ordinary_paragraph_pipes_are_untouched(self):
        document = Document()
        document.add_paragraph("a | b is not a table")

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertEqual(markdown, "a | b is not a table")

    def test_table_with_empty_cells_still_emits_a_delimiter_row(self):
        document = _build_table(1, 2, [["", ""]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertIn("| --- | --- |", markdown)

    def test_document_without_tables_is_unchanged(self):
        document = Document()
        document.add_heading("Doc Title", level=1)
        document.add_paragraph("Body text")

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertEqual(markdown, "# Doc Title\n\nBody text")


class TestDocxImportKeepsDocumentOrder(unittest.TestCase):
    """Tables and paragraphs have to be exported the way the document reads.

    ``document.paragraphs`` and ``document.tables`` are two independent lists, so
    exporting one and then the other moves every table to the end of the output
    regardless of where it sat. A report whose summary table belongs between two
    paragraphs used to arrive with the prose first and the table last, which
    silently rewrites what the document said.
    """

    @staticmethod
    def _add_table(document, values):
        table = document.add_table(rows=len(values), cols=len(values[0]))
        table.style = "Table Grid"
        for row_index, row in enumerate(values):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
        return table

    def test_text_after_a_table_stays_after_it(self):
        document = Document()
        document.add_paragraph("Introduction.")
        self._add_table(document, [["Metric", "Value"], ["Latency", "42ms"]])
        document.add_paragraph("Conclusion.")

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertLess(
            markdown.index("Introduction."),
            markdown.index("| Metric | Value |"),
        )
        self.assertLess(
            markdown.index("| Metric | Value |"),
            markdown.index("Conclusion."),
        )

    def test_text_before_a_table_stays_before_it(self):
        document = Document()
        document.add_paragraph("Setup follows.")
        self._add_table(document, [["A", "B"]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertLess(markdown.index("Setup follows."), markdown.index("| A | B |"))

    def test_two_tables_around_a_paragraph_are_not_swapped(self):
        document = Document()
        self._add_table(document, [["first-header", "x"]])
        document.add_paragraph("between the tables")
        self._add_table(document, [["second-header", "y"]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertLess(
            markdown.index("first-header"),
            markdown.index("between the tables"),
        )
        self.assertLess(
            markdown.index("between the tables"),
            markdown.index("second-header"),
        )

    def test_consecutive_tables_keep_their_own_delimiter_rows(self):
        document = Document()
        self._add_table(document, [["one-a", "one-b"], ["1", "2"]])
        self._add_table(document, [["two-a", "two-b"], ["3", "4"]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertEqual(markdown.count("| --- | --- |"), 2)
        self.assertLess(markdown.index("one-a"), markdown.index("two-a"))
        self.assertIn("| 1 | 2 |", markdown)
        self.assertIn("| 3 | 4 |", markdown)

    def test_a_heading_before_a_table_keeps_its_lead(self):
        document = Document()
        document.add_heading("Results", level=2)
        self._add_table(document, [["Score", "10"]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertLess(markdown.index("## Results"), markdown.index("| Score | 10 |"))

    def test_a_table_at_the_very_start_is_not_pushed_down(self):
        document = Document()
        self._add_table(document, [["Leading", "1"]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertTrue(markdown.startswith("| Leading | 1 |"))

    def test_a_table_at_the_very_end_stays_at_the_very_end(self):
        document = Document()
        document.add_paragraph("Trailing prose.")
        self._add_table(document, [["Final", "9"]])

        markdown = _convert_docx_to_markdown(_save(document))

        lines = markdown.splitlines()
        self.assertEqual(lines[-2:], ["| Final | 9 |", "| --- | --- |"])
        self.assertLess(
            markdown.index("Trailing prose."), markdown.index("| Final | 9 |")
        )

    def test_a_full_interleaved_document_matches_the_body_order_exactly(self):
        document = Document()
        document.add_paragraph("Intro")
        self._add_table(document, [["A", "B"], ["1", "2"]])
        document.add_paragraph("Middle")
        document.add_heading("Notes", level=1)
        document.add_paragraph("Outro")

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertEqual(
            markdown,
            "Intro\n\n"
            "| A | B |\n"
            "| --- | --- |\n"
            "| 1 | 2 |\n\n"
            "Middle\n\n"
            "# Notes\n\n"
            "Outro",
        )

    def test_list_items_around_a_table_keep_their_markers_and_order(self):
        document = Document()
        document.add_paragraph("First bullet", style="List Bullet")
        self._add_table(document, [["Bullet", "Point"]])
        document.add_paragraph("Second bullet", style="List Bullet")

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertLess(
            markdown.index("- First bullet"), markdown.index("| Bullet | Point |")
        )
        self.assertLess(
            markdown.index("| Bullet | Point |"), markdown.index("- Second bullet")
        )

    def test_an_empty_paragraph_between_blocks_does_not_reorder_them(self):
        document = Document()
        document.add_paragraph("Before")
        document.add_paragraph("")  # dropped, but must not shift the table
        self._add_table(document, [["Middle", "7"]])

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertEqual(markdown, "Before\n\n| Middle | 7 |\n| --- | --- |")

    def test_a_table_nested_in_a_cell_is_still_not_emitted(self):
        # Only top-level blocks were exported before, and that set must not grow.
        document = Document()
        document.add_paragraph("Host text")
        outer = self._add_table(document, [["Host", "Cell"]])
        inner = outer.cell(0, 1).add_table(rows=1, cols=1)
        inner.cell(0, 0).text = "nested-only"
        document.add_paragraph("After text")

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertNotIn("nested-only", markdown)
        self.assertLess(markdown.index("Host text"), markdown.index("| Host | Cell |"))
        self.assertLess(markdown.index("| Host | Cell |"), markdown.index("After text"))

    def test_a_revision_tracked_paragraph_is_still_skipped(self):
        # python-docx omits w:ins content from both .paragraphs and
        # iter_inner_content; the table must not leak into that gap.
        document = Document()
        document.add_paragraph("Committed before")
        tracked = OxmlElement("w:ins")
        para = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "tracked insert"
        run.append(text)
        para.append(run)
        tracked.append(para)
        document.element.body.insert(1, tracked)
        self._add_table(document, [["After", "tracked"]])
        document.add_paragraph("Committed after")

        markdown = _convert_docx_to_markdown(_save(document))

        self.assertNotIn("tracked insert", markdown)
        self.assertLess(
            markdown.index("Committed before"), markdown.index("| After | tracked |")
        )
        self.assertLess(
            markdown.index("| After | tracked |"), markdown.index("Committed after")
        )

    def test_ordering_survives_the_renderer(self):
        document = Document()
        document.add_paragraph("Above the table.")
        self._add_table(document, [["Metric", "Value"], ["Latency", "42ms"]])
        document.add_paragraph("Below the table.")

        html = render_markdown(_convert_docx_to_markdown(_save(document)))

        self.assertIn("<table>", html)
        self.assertLess(html.index("Above the table."), html.index("<table>"))
        self.assertLess(html.index("</table>"), html.index("Below the table."))

    def test_table_helper_returns_rows_without_a_trailing_blank(self):
        document = Document()
        table = self._add_table(document, [["A", "B"], ["1", "2"]])

        from backend.routers.files import _docx_table_to_markdown

        self.assertEqual(
            _docx_table_to_markdown(table),
            ["| A | B |", "| --- | --- |", "| 1 | 2 |"],
        )

    def test_table_helper_is_empty_for_a_table_with_no_rows(self):
        from backend.routers.files import _docx_table_to_markdown

        document = Document()
        table = document.add_table(rows=0, cols=2)
        table._tbl.remove(table.rows[0]._tr if table.rows else table._tbl[0])

        self.assertEqual(_docx_table_to_markdown(table), [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
